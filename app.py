import streamlit as st
from google import genai
from google.genai import types
import requests
from bs4 import BeautifulSoup
import pypdf
import datetime
import json
import os
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt

# 페이지 설정
st.set_page_config(page_title="과학 탐구 수업 & 교-수-평-기 설계 비서", page_icon="🧪", layout="wide")

# ==========================================
# 📁 노트북 저장 폴더
# ==========================================
BASE_DIR = Path(__file__).resolve().parent
CHAT_DIR = BASE_DIR / "저장된_대화"
OUTPUT_DIR = BASE_DIR / "결과물"
SECRETS_PATH = BASE_DIR / ".streamlit" / "secrets.toml"
CHAT_DIR.mkdir(exist_ok=True)
OUTPUT_DIR.mkdir(exist_ok=True)

DEFAULT_MODEL = "gemini-flash-latest"


def safe_filename(text):
    return re.sub(r'[\\/:*?"<>|\n\r\t]', "_", text).strip()[:40] or "제목없음"


def save_chat(chat_id, chat):
    with open(CHAT_DIR / f"{chat_id}.json", "w", encoding="utf-8") as f:
        json.dump(chat, f, ensure_ascii=False, indent=2)


def load_chats():
    chats = {}
    for path in sorted(CHAT_DIR.glob("*.json")):
        try:
            with open(path, encoding="utf-8") as f:
                chats[path.stem] = json.load(f)
        except Exception:
            pass
    return chats


def new_chat_id():
    return datetime.datetime.now().strftime("%Y%m%d_%H%M%S_%f")


# ==========================================
# 📄 마크다운 → Word(.docx) 변환
# ==========================================
def add_inline(paragraph, text):
    # **굵게** 처리
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part:
            paragraph.add_run(part).bold = (i % 2 == 1)


def markdown_to_docx(md_text, title, path):
    doc = Document()
    doc.styles["Normal"].font.name = "맑은 고딕"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading(title, level=0)

    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        # 표 (| a | b |)
        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            if rows:
                ncols = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=ncols)
                table.style = "Table Grid"
                for r, row in enumerate(rows):
                    for c, cell in enumerate(row):
                        p = table.cell(r, c).paragraphs[0]
                        add_inline(p, cell)
                        if r == 0:
                            for run in p.runs:
                                run.bold = True
            continue

        heading = re.match(r"^(#{1,6})\s+(.*)", stripped)
        bullet = re.match(r"^[-*+]\s+(.*)", stripped)
        number = re.match(r"^\d+[.)]\s+(.*)", stripped)
        if not stripped:
            pass
        elif re.fullmatch(r"[-*_]{3,}", stripped):
            doc.add_paragraph("─" * 30)
        elif heading:
            doc.add_heading(heading.group(2).replace("**", ""), level=min(len(heading.group(1)), 4))
        elif bullet:
            add_inline(doc.add_paragraph(style="List Bullet"), bullet.group(1))
        elif number:
            add_inline(doc.add_paragraph(style="List Number"), number.group(1))
        else:
            add_inline(doc.add_paragraph(), stripped)
        i += 1

    doc.save(path)


# ==========================================
# 🔄 외부 데이터 캐싱 (속도 최적화)
# ==========================================
@st.cache_data(ttl=86400)
def fetch_guide_data():
    try:
        url = "https://ppzine.kr/class_edu/"
        headers = {"User-Agent": "Mozilla/5.0"}
        res = requests.get(url, headers=headers, timeout=3)
        soup = BeautifulSoup(res.text, 'html.parser')
        texts = [el.get_text(strip=True) for el in soup.select('p, div, td') if len(el.get_text(strip=True)) > 15]
        return "\n".join(texts[:3]) if texts else "2022 개정 깊이있는 학습 지침"
    except Exception:
        return "2022 개정 깊이있는 학습 및 질문 중심 수업 지침"

@st.cache_data(ttl=86400)
def fetch_steam_data():
    try:
        session = requests.Session()
        headers = {"User-Agent": "Mozilla/5.0", "Referer": "https://steam.kosac.re.kr/"}
        main_url = "https://steam.kosac.re.kr/learning/curriculum/list/menu/220"
        session.get(main_url, headers=headers, timeout=3)
        payload = {'schulClCode': 'M', 'curriculumCd': '220', 'subjcCode': 'SCI', 'pageIndex': '1'}
        res = session.post(main_url, headers=headers, data=payload, timeout=3)
        soup = BeautifulSoup(res.text, 'html.parser')
        titles = [row.get_text(strip=True).split('\n')[0] for row in soup.select('tbody tr, .subject, td.al') if len(row.get_text(strip=True)) > 5]
        return "\n".join(titles[:3]) if titles else "STEAM 중등 과학 융합 탐구 모델"
    except Exception:
        return "중학교 과학 교과중심 STEAM 교수학습자료"


@st.cache_data(ttl=3600, show_spinner=False)
def list_models(api_key):
    """이 API 키로 사용 가능한 Gemini 모델 목록 (서비스 종료된 모델은 자동 제외)"""
    try:
        client = genai.Client(api_key=api_key)
        names = []
        for m in client.models.list():
            if "generateContent" in (m.supported_actions or []) and "gemini" in m.name:
                names.append(m.name.replace("models/", ""))
        return sorted(names)
    except Exception:
        return []


# ==========================================
# 💾 세션 상태 및 채팅 히스토리 관리 (노트북 파일에서 불러오기)
# ==========================================
if "chat_history_list" not in st.session_state:
    st.session_state.chat_history_list = load_chats()  # { chat_id: {"title": ..., "messages": [...] } }

if "current_chat_id" not in st.session_state:
    st.session_state.current_chat_id = new_chat_id()
    st.session_state.chat_history_list[st.session_state.current_chat_id] = {
        "title": "새 대화",
        "messages": []
    }

# 사이드바 설정
st.sidebar.title("⚙️ 설정 및 외부 자료")

try:
    saved_key = st.secrets.get("GEMINI_API_KEY", "")
    saved_model = st.secrets.get("GEMINI_MODEL", DEFAULT_MODEL)
except Exception:
    saved_key, saved_model = "", DEFAULT_MODEL

if saved_key:
    api_key = saved_key
    st.sidebar.success("🔑 저장된 API Key 사용 중")
else:
    api_key = st.sidebar.text_input("Gemini API Key 입력", type="password")
    if api_key and st.sidebar.button("💾 이 노트북에 Key 저장 (다음부터 입력 생략)"):
        SECRETS_PATH.parent.mkdir(exist_ok=True)
        with open(SECRETS_PATH, "w", encoding="utf-8") as f:
            f.write(f'GEMINI_API_KEY = "{api_key.strip()}"\n')
        st.sidebar.success("저장했습니다. 다음 실행부터 자동으로 사용됩니다.")

model_name = saved_model
if api_key:
    available = list_models(api_key)
    options = [DEFAULT_MODEL] + [m for m in available if m != DEFAULT_MODEL]
    default_index = options.index(saved_model) if saved_model in options else 0
    model_name = st.sidebar.selectbox("사용할 Gemini 모델", options, index=default_index,
                                      help="gemini-flash-latest는 항상 최신 Flash 모델을 가리킵니다.")

# ➕ 새 채팅 버튼
if st.sidebar.button("➕ 새 대화 시작하기", use_container_width=True):
    new_id = new_chat_id()
    st.session_state.current_chat_id = new_id
    st.session_state.chat_history_list[new_id] = {
        "title": "새 대화",
        "messages": []
    }
    st.rerun()

# 🗂️ 지난 대화 목록 선택 (제목이 같아도 ID로 구분, 최신순)
st.sidebar.markdown("---")
st.sidebar.subheader("🗂️ 대화 기록 목록")

chat_ids = sorted(st.session_state.chat_history_list.keys(), reverse=True)

if chat_ids:
    # 제목 뒤에 날짜·시간을 붙여 같은 제목도 구분
    labels = {
        f"{st.session_state.chat_history_list[cid]['title']} ({cid[4:6]}/{cid[6:8]} {cid[9:11]}:{cid[11:13]}:{cid[13:15]})": cid
        for cid in chat_ids
    }
    label_list = list(labels.keys())
    selected_index = chat_ids.index(st.session_state.current_chat_id) if st.session_state.current_chat_id in chat_ids else 0
    selected_label = st.sidebar.radio(
        "저장된 대화 선택",
        options=label_list,
        index=selected_index,
        label_visibility="collapsed"
    )
    st.session_state.current_chat_id = labels[selected_label]

if st.sidebar.button("📂 결과물 폴더 열기", use_container_width=True):
    if hasattr(os, "startfile"):
        os.startfile(OUTPUT_DIR)
    else:
        st.sidebar.info(f"결과물 폴더: {OUTPUT_DIR}")

# 📚 교과서 PDF 업로드
st.sidebar.markdown("---")
st.sidebar.subheader("📚 교과서 PDF 업로드")
uploaded_pdf = st.sidebar.file_uploader("교과서 단원 PDF 파일 선택", type=["pdf"])

pdf_text_content = ""
if uploaded_pdf is not None:
    try:
        pdf_reader = pypdf.PdfReader(uploaded_pdf)
        extracted_pages = [page.extract_text() for page in pdf_reader.pages if page.extract_text()]
        pdf_text_content = "\n".join(extracted_pages)
        st.sidebar.success(f"교과서 PDF 학습 완료! ({len(pdf_reader.pages)}페이지)")
    except Exception:
        st.sidebar.error("PDF 읽기 중 오류가 발생했습니다.")

guide_context = fetch_guide_data()
steam_context = fetch_steam_data()

st.sidebar.markdown("---")
st.sidebar.subheader("🌐 외부 데이터 자동 수집 상태")
st.sidebar.success("✅ 깊이있는 학습 지침서 연동")
st.sidebar.success("✅ STEAM 중등 과학 포털 연동")

# 메인 화면
st.title("🧪 중학교 과학 질문 중심 수업 설계 비서")
st.caption("교과서 중심 | 2022 개정 깊이 있는 학습 | 질문 중심 Scaffolding | 교-수-평-기 일체화")

if not api_key:
    st.info("👈 왼쪽 사이드바에 Gemini API Key를 입력하세요.")
    st.stop()

SYSTEM_PROMPT = """
너는 중학교 과학 교사를 돕는 "상호작용형 질문 중심 수업 설계 및 교-수-평-기 일체화 AI 비서"이다.

[교과서 기반 학습 최우선 지침]
1. 업로드된 교과서 PDF 내용이 존재할 경우, 해당 교과서에 기술된 실제 개념 설명, 탐구 활동, 용어, 소단원 체계를 중심(Core)으로 삼는다.
2. 교과서 수록 질문이나 자료를 활용하여 학생용 비계(Scaffolding) 및 발문을 설계한다.

[2022 개정 교육과정: 깊이 있는 학습 반영 지침]
1. 개념 기반 탐구(Conceptual Inquiry)와 핵심 아이디어 중심의 이해를 유도한다.
2. 학생이 핵심 개념 속에서 '삶과 연계된 핵심 질문'을 발견하도록 비계를 설계한다.

[상호작용 단계]
1단계: 교과서 단원/내용이 들어오면, 이 단원에 가장 자연스러운 탐구 소스 2~3가지를 제안하고 교사의 선택을 묻는다.
2단계: 교사의 선택에 맞춰 교실 준비물, 차시 제약 등을 질문하여 맞춘다.
3단계: [교-수-평-기 일체화 패키지] (활동지, 교사 가이드, 루브릭, 세특 예시문)를 완성한다.
"""

# 참고 자료는 시스템 지시에 넣어 대화 내내 유지
system_instruction = SYSTEM_PROMPT
if pdf_text_content:
    system_instruction += f"\n\n[학습된 교과서 PDF 내용]:\n{pdf_text_content[:3000]}"
if guide_context:
    system_instruction += f"\n\n[참고 지침서]: {guide_context[:500]}"
if steam_context:
    system_instruction += f"\n\n[참고 STEAM]: {steam_context[:500]}"

# 현재 활성화된 채팅 메시지 가져오기
current_chat_id = st.session_state.current_chat_id
current_chat = st.session_state.chat_history_list[current_chat_id]
current_messages = current_chat["messages"]

# 대화 내용 표시 (AI 답변마다 Word 저장 버튼)
for idx, msg in enumerate(current_messages):
    with st.chat_message(msg["role"]):
        st.markdown(msg["content"])
        if msg["role"] == "assistant":
            if st.button("📄 Word(.docx)로 저장", key=f"docx_{current_chat_id}_{idx}"):
                title = current_chat["title"].rstrip(".")
                filename = f"{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}_{safe_filename(title)}.docx"
                try:
                    markdown_to_docx(msg["content"], title, OUTPUT_DIR / filename)
                    st.success(f"저장 완료: 결과물\\{filename}")
                except Exception as e:
                    st.error(f"저장 중 오류: {e} (같은 파일이 Word에서 열려 있으면 닫고 다시 시도하세요)")

if prompt := st.chat_input("교과서 단원명이나 다루고 싶은 주제를 입력하세요..."):
    # 첫 질문일 경우 채팅 제목 변경
    if len(current_messages) == 0:
        current_chat["title"] = prompt[:15] + ("..." if len(prompt) > 15 else "")

    current_messages.append({"role": "user", "content": prompt})
    save_chat(current_chat_id, current_chat)
    with st.chat_message("user"):
        st.markdown(prompt)

    # 이전 대화 전체를 함께 보내 1→2→3단계 흐름 유지
    contents = [
        types.Content(role="user" if m["role"] == "user" else "model", parts=[types.Part(text=m["content"])])
        for m in current_messages
    ]

    with st.chat_message("assistant"):
        with st.spinner("⚡ 수업을 구성하는 중입니다..."):
            try:
                client = genai.Client(api_key=api_key)
                response = client.models.generate_content(
                    model=model_name,
                    contents=contents,
                    config=types.GenerateContentConfig(system_instruction=system_instruction),
                )

                st.markdown(response.text)
                current_messages.append({"role": "assistant", "content": response.text})
                save_chat(current_chat_id, current_chat)
                st.rerun()
            except Exception as e:
                st.error(f"API 호출 오류 발생: {e}\n\n👉 API Key와 사이드바의 모델 선택을 확인해 주세요!")
